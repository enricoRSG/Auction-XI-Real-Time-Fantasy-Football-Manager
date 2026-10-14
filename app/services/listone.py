from __future__ import annotations

from io import BytesIO
from openpyxl import load_workbook


# The app supports both:
# 1) the older/custom Leghe FC export used during development
# 2) the current official Fantacalcio quotation workbook (2026/27), whose
#    first row is a title and whose real headers are on the second row.
HEADER_ALIASES = {
    "source_id": ("#", "Id"),
    "name": ("Nome",),
    "serie_a_team": ("Sq.", "Squadra"),
    "role": ("R.", "R"),
    "mantra_role": ("R.MANTRA", "RM"),
    "fvm": ("FVM/1000", "FVM"),
    # Qt.A = current quotation; QUOT. is the legacy/current quotation field
    "quotation": ("QUOT.", "Qt.A"),
    "out_of_list": ("Fuori lista",),
}

REQUIRED_FIELDS = {"name", "serie_a_team", "role", "fvm"}


def parse_listone_xlsx(content: bytes) -> list[dict]:
    wb = load_workbook(BytesIO(content), data_only=True, read_only=True)

    # Official Fantacalcio files expose a "Tutti" sheet. Prefer it so we do
    # not accidentally import the separate "Ceduti" sheet.
    ws = wb["Tutti"] if "Tutti" in wb.sheetnames else wb[wb.sheetnames[0]]

    header_map, header_row_number = _find_header_map(ws)
    missing = REQUIRED_FIELDS - set(header_map)
    if missing:
        readable = ", ".join(sorted(missing))
        raise ValueError(f"Colonne mancanti nel listone: {readable}")

    result: list[dict] = []
    for row in ws.iter_rows(min_row=header_row_number + 1, values_only=True):
        name = _value(row, header_map.get("name"))
        if name in (None, ""):
            continue

        role = str(_value(row, header_map.get("role")) or "").strip().upper()
        # Only Classic roles are valid in the current V1.
        if role not in {"P", "D", "C", "A"}:
            continue

        out_value = _value(row, header_map.get("out_of_list"))

        result.append(
            {
                "source_id": _to_int(_value(row, header_map.get("source_id"))),
                "name": str(name).strip(),
                "serie_a_team": str(_value(row, header_map.get("serie_a_team")) or "").strip(),
                "role": role,
                "mantra_role": _optional_text(_value(row, header_map.get("mantra_role"))),
                "fvm": _to_int(_value(row, header_map.get("fvm"))),
                "quotation": _to_int(_value(row, header_map.get("quotation"))),
                "out_of_list": _to_bool(out_value),
            }
        )

    if not result:
        raise ValueError("Il file non contiene calciatori validi.")
    return result


def _find_header_map(ws) -> tuple[dict[str, int], int]:
    """Find the real header row and normalize supported column names.

    Current official Fantacalcio workbooks have a title in row 1 and headers
    in row 2, while the legacy file has headers in row 1. Scanning the first
    10 rows keeps the importer resilient to small layout changes.
    """
    alias_to_field = {
        alias.casefold(): field
        for field, aliases in HEADER_ALIASES.items()
        for alias in aliases
    }

    best_map: dict[str, int] = {}
    best_row = 1

    for row_number, row in enumerate(
        ws.iter_rows(min_row=1, max_row=min(ws.max_row, 10), values_only=True),
        start=1,
    ):
        candidate: dict[str, int] = {}
        for idx, value in enumerate(row):
            if value is None:
                continue
            key = str(value).strip().casefold()
            field = alias_to_field.get(key)
            if field and field not in candidate:
                candidate[field] = idx

        if len(candidate) > len(best_map):
            best_map = candidate
            best_row = row_number

        if REQUIRED_FIELDS.issubset(candidate):
            return candidate, row_number

    return best_map, best_row


def _value(row, idx):
    if idx is None or idx >= len(row):
        return None
    return row[idx]


def _optional_text(value):
    if value in (None, ""):
        return None
    return str(value).strip()


def _to_bool(value):
    if value in (None, "", 0, False):
        return False
    if isinstance(value, str):
        return value.strip().casefold() not in {"", "0", "false", "no", "n"}
    return bool(value)


def _to_int(value):
    if value in (None, ""):
        return None
    try:
        return int(round(float(value)))
    except (TypeError, ValueError):
        return None
