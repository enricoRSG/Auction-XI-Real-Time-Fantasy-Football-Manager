from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import select

from app.database import Base, DB_PATH, SessionLocal, engine
from app.main import app
from app.models import AuctionPlayer, Player

LISTONE = Path(__file__).resolve().parent / "fixtures" / "listone_full_simulation.xlsx"


def setup_module():
    engine.dispose()
    if DB_PATH.exists():
        DB_PATH.unlink()
    Base.metadata.create_all(bind=engine)


def test_complete_10_team_auction_with_patching_exports_and_reset(tmp_path):
    """Simulate the full agreed league shape: 10 teams, 500 credits, 3/8/8/6."""
    client = TestClient(app)
    team_names = "\n".join(f"Team {i}" for i in range(1, 11))
    with LISTONE.open("rb") as fh:
        setup = client.post(
            "/api/setup",
            data={
                "name": "Full league simulation",
                "auctioneer_type": "participant",
                "team_count": "10",
                "starting_credits": "500",
                "slots_p": "3",
                "slots_d": "8",
                "slots_c": "8",
                "slots_a": "6",
                "release_p": "1",
                "release_d": "2",
                "release_c": "2",
                "release_a": "2",
                "patch_limit": "3",
                "base_rule": "fixed",
                "base_value": "1",
                "allow_overspend": "true",
                "call_mode": "manual",
                "team_names": team_names,
                "auctioneer_team_index": "0",
            },
            files={"listone": (LISTONE.name, fh, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
    assert setup.status_code == 200, setup.text
    body = setup.json()
    auction_id = body["auction_id"]
    admin_token = body["admin_url"].rsplit("/", 1)[-1]

    links_r = client.get(f"/api/auction/{auction_id}/admin/access-links", params={"admin_token": admin_token})
    assert links_r.status_code == 200
    links = links_r.json()
    assert len(links) == 10

    with SessionLocal() as db:
        by_role = {}
        for role in "PDCA":
            rows = db.execute(
                select(Player.id)
                .join(AuctionPlayer, AuctionPlayer.player_id == Player.id)
                .where(AuctionPlayer.auction_id == auction_id, Player.role == role, AuctionPlayer.state == "SVINCOLATO")
                .order_by(Player.id)
            ).scalars().all()
            by_role[role] = list(rows)

    required = {"P": 3, "D": 8, "C": 8, "A": 6}
    cursors = {r: 0 for r in required}

    # Fill every roster with 250 official acquisitions at 1 credit each.
    for team in links:
        team_token = team["url"].rsplit("/", 1)[-1]
        for role, count in required.items():
            for _ in range(count):
                player_id = by_role[role][cursors[role]]
                cursors[role] += 1
                start = client.post(
                    f"/api/auction/{auction_id}/start-player",
                    params={"admin_token": admin_token},
                    json={"player_id": player_id},
                )
                assert start.status_code == 200, start.text
                bid = client.post(
                    f"/api/auction/{auction_id}/bid",
                    json={"team_id": team["id"], "amount": 1, "access_token": team_token},
                )
                assert bid.status_code == 200, bid.text
                award = client.post(f"/api/auction/{auction_id}/award", params={"admin_token": admin_token})
                assert award.status_code == 200, award.text
                assert award.json()["unsold"] is False

    # All teams have complete rosters and 475 credits left.
    teams_r = client.get(f"/api/auction/{auction_id}/teams", params={"admin_token": admin_token})
    assert teams_r.status_code == 200
    for team in teams_r.json():
        assert team["credits_remaining"] == 475
        assert team["counts"] == required

    concluded = client.post(f"/api/auction/{auction_id}/conclude-live", params={"admin_token": admin_token})
    assert concluded.status_code == 200, concluded.text
    assert concluded.json()["status"] == "RIEPILOGO_SVINCOLATI"

    free_ranked = client.get(f"/api/auction/{auction_id}/free-ranked").json()
    assert len(free_ranked) == 4

    started_patch = client.post(f"/api/auction/{auction_id}/start-patching", params={"admin_token": admin_token})
    assert started_patch.status_code == 200

    team = links[0]
    roster = client.get(f"/api/auction/{auction_id}/team/{team['id']}/roster",params={'admin_token':admin_token}).json()
    released_player = roster["players"]["A"][0]
    patch_release_r = client.post(
        f"/api/auction/{auction_id}/patch-release",
        params={"admin_token": admin_token},
        json={"team_id": team["id"], "player_id": released_player["id"]},
    )
    assert patch_release_r.status_code == 200, patch_release_r.text
    move = patch_release_r.json()

    free_attackers = client.get(f"/api/auction/{auction_id}/free-ranked", params={"role": "A"}).json()
    replacement = next(p for p in free_attackers if p["id"] != released_player["id"])
    patch_assign_r = client.post(
        f"/api/auction/{auction_id}/patch-assign",
        params={"admin_token": admin_token},
        json={"team_id": team["id"], "move_id": move["move_id"], "player_id": replacement["id"], "price": 1},
    )
    assert patch_assign_r.status_code == 200, patch_assign_r.text

    # The released player is permanently banned for the same team.
    # It remains free, but cannot be the replacement in a later move by that team.
    final_patch = client.post(f"/api/auction/{auction_id}/conclude-patching", params={"admin_token": admin_token})
    assert final_patch.status_code == 200, final_patch.text
    assert final_patch.json()["status"] == "ROSE_DEFINITIVE"

    csv_r = client.get(f"/api/auction/{auction_id}/export/leghe", params={"admin_token": admin_token})
    assert csv_r.status_code == 200
    lines = csv_r.content.decode("utf-8-sig").strip().splitlines()
    assert len(lines) == 251  # header + 250 final players

    xlsx_r = client.get(f"/api/auction/{auction_id}/export/archive", params={"admin_token": admin_token})
    assert xlsx_r.status_code == 200
    export_path = tmp_path / 'full_sim_archive.xlsx'
    export_path.write_bytes(xlsx_r.content)
    wb = load_workbook(export_path, data_only=True)
    assert set(wb.sheetnames) == {"Rose", "Crediti"}
    assert wb["Rose"].max_row == 251
    assert wb["Crediti"].max_row == 11

    reset = client.post(
        f"/api/auction/{auction_id}/close-and-reset",
        params={"admin_token": admin_token, "exports_confirmed": "true"},
    )
    assert reset.status_code == 200, reset.text
    assert reset.json()["redirect"] == "/"


def test_setup_rejects_duplicate_team_names_and_bad_auctioneer_index():
    client = TestClient(app)
    with LISTONE.open("rb") as fh:
        duplicate = client.post(
            "/api/setup",
            data={
                "auctioneer_type": "participant", "team_count": "2", "starting_credits": "500",
                "slots_p": "1", "slots_d": "0", "slots_c": "0", "slots_a": "0",
                "release_p": "0", "release_d": "0", "release_c": "0", "release_a": "0",
                "patch_limit": "0", "base_rule": "fixed", "base_value": "1",
                "allow_overspend": "false", "call_mode": "manual",
                "team_names": "Uguale\nuguale", "auctioneer_team_index": "0",
            },
            files={"listone": (LISTONE.name, fh, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
    assert duplicate.status_code == 400
    assert "univoci" in duplicate.json()["detail"]

    with LISTONE.open("rb") as fh:
        bad_index = client.post(
            "/api/setup",
            data={
                "auctioneer_type": "participant", "team_count": "2", "starting_credits": "500",
                "slots_p": "1", "slots_d": "0", "slots_c": "0", "slots_a": "0",
                "release_p": "0", "release_d": "0", "release_c": "0", "release_a": "0",
                "patch_limit": "0", "base_rule": "fixed", "base_value": "1",
                "allow_overspend": "false", "call_mode": "manual",
                "team_names": "A\nB", "auctioneer_team_index": "5",
            },
            files={"listone": (LISTONE.name, fh, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
    assert bad_index.status_code == 400
    assert "squadra battitore valida" in bad_index.json()["detail"]
